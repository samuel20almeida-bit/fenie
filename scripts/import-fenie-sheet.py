"""Read Fenié's CATALOGO worksheet and emit a commercial-only staging payload.

Usage: python3 scripts/import-fenie-sheet.py SOURCE.xlsx DESTINATION_DIRECTORY
Requires openpyxl. Source workbook is read-only and never copied to the repo.
"""
import collections
import json
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

# Explicit allowlist. Costs, margins, price floors and marketplace data are not read.
SOURCE_COLUMNS = {
    'brand': 1, 'line': 2, 'sku': 3, 'name': 4,
    'price': 12, 'situation': 18,
}
SIZE_PATTERN = re.compile(
    r'(?<![\w.])(\d+(?:[.,]\d+)?)\s*(ml|kg|g|lt|l|litros?|un(?:id(?:ades?)?)?|folhas?)\b',
    re.I,
)


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def slugify(value):
    ascii_value = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', ascii_value.lower()).strip('-')


def cents(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        price = Decimal(str(value).replace(',', '.'))
        if not price.is_finite() or price <= 0:
            return None
        rounded = price.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
        if price != rounded:
            return None  # Never silently change a source commercial price.
        result = int(rounded * 100)
        return result if result <= 100000000 else None
    except (InvalidOperation, ValueError, TypeError):
        return None


def size_from_name(name):
    matches = []
    for match in SIZE_PATTERN.finditer(name):
        number, unit = match.groups()
        unit = unit.lower()
        if unit in ['lt', 'l', 'litro', 'litros']:
            unit = 'L'
        elif unit.startswith('un'):
            unit = 'un.'
        elif unit.startswith('folha'):
            unit = 'folhas'
        value = f'{number} {unit}'
        if value not in matches:
            matches.append(value)
    return ' + '.join(matches)


def extract_rows(worksheet):
    result = []
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
        if row_number <= 2:
            continue
        source = {key: row[index] if index < len(row) else None for key, index in SOURCE_COLUMNS.items()}
        if not source['sku'] and not source['name']:
            continue
        source['sourceRow'] = row_number
        result.append(source)
    return result


def prepare(rows):
    sku_counts = collections.Counter(clean(r['sku']) for r in rows if clean(r['sku']))
    products, review = [], []
    used_slugs = set()
    for row in rows:
        sku, name, brand, line = (clean(row[key]) for key in ['sku', 'name', 'brand', 'line'])
        price_cents = cents(row['price'])
        size = size_from_name(name)
        blockers = []
        if not sku:
            blockers.append('SKU ausente')
        elif sku_counts[sku] > 1:
            blockers.append('SKU duplicado: corrigir na fonte')
        if not name:
            blockers.append('Nome ausente')
        if not brand:
            blockers.append('Marca ausente')
        if not line:
            blockers.append('Linha ausente')
        if price_cents is None:
            blockers.append('Preço praticado ausente ou inválido')
        if not size:
            blockers.append('Embalagem/volume não explícito na descrição')
        if clean(row['situation']).lower() != 'ativo':
            blockers.append('Situação comercial não confirmada como Ativo')
        slug = slugify(f'{name}-{sku}')
        if not slug or slug in used_slugs:
            blockers.append('Slug ausente ou duplicado')
        if not blockers:
            used_slugs.add(slug)
            products.append({
                'id': sku, 'sku': sku, 'slug': slug, 'name': name,
                'brand': brand, 'line': line, 'category': line,
                'size': size, 'priceCents': price_cents,
                'status': 'consult', 'order': len(products),
                'keywords': [sku, brand, line],
                'reviewPending': ['Foto oficial', 'Taxonomia final', 'Revisão comercial'],
            })
        review.append({
            'sourceRow': row['sourceRow'], 'sku': sku, 'name': name,
            'brand': brand, 'line': line, 'size': size,
            'priceCents': price_cents, 'situation': clean(row['situation']),
            'eligibleForPreview': not blockers, 'blockers': blockers,
        })
    counts = collections.Counter(issue for r in review for issue in r['blockers'])
    summary = {
        'sourceRows': len(rows), 'uniqueSkus': len(sku_counts),
        'previewProducts': len(products), 'reviewRows': sum(bool(r['blockers']) for r in review),
        'priceField': 'CATALOGO!M — PREÇO PRATICADO',
        'brandCounts': dict(collections.Counter(clean(r['brand']) or 'Sem marca' for r in rows)),
        'blockerCounts': dict(counts),
        'allProductsMissingOfficialImages': True,
        'commercialOrdersEnabled': False,
    }
    payload = {'demo': False, 'reviewOnly': True, 'taxonomyMode': 'source-line', 'summary': summary, 'products': products}
    return payload, review


def main():
    import openpyxl
    if len(sys.argv) != 3:
        raise SystemExit('Uso: python3 scripts/import-fenie-sheet.py SOURCE.xlsx DESTINATION_DIRECTORY')
    workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    sheet = workbook['CATALOGO']
    headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    expected = {1: 'MARCA', 2: 'LINHA', 3: 'CÓDIGO', 4: 'DESCRIÇÃO', 12: 'PREÇO PRATICADO', 18: 'Situação'}
    if any(index >= len(headers) or clean(headers[index]) != name for index, name in expected.items()):
        raise SystemExit('Cabeçalho mudou: interrompendo importação para não escolher os campos errados.')
    payload, review = prepare(extract_rows(sheet))
    destination = Path(sys.argv[2]); destination.mkdir(parents=True, exist_ok=True)
    for name, content in [('catalog-preview.json', payload), ('catalog-review.json', review)]:
        (destination / name).write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(payload['summary'], ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
