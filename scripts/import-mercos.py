"""Read a Mercos export; publish only commercial fields and the standard price.

Usage: python3 scripts/import-mercos.py SOURCE.xlsx PREVIOUS.json OUTPUT.json
The export and reconciliation report must stay outside the public repository.
"""
import collections
import importlib.util
import json
import sys
from pathlib import Path

spec = importlib.util.spec_from_file_location('sheet', Path(__file__).with_name('import-fenie-sheet.py'))
sheet = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sheet)

BRANDS = {'OLENKA', 'GLYNETT', 'MUP COLOR', 'PROHALL', 'RIGOLIM', 'DOHA', 'SPA DO FIO', 'MUP MAKEUP'}
COLUMNS = {'sku': 0, 'name': 1, 'price': 2, 'unit': 8, 'group': 16, 'category': 17, 'subcategory': 18, 'inactive': 19, 'hidden': 20}
HEADERS = {0: 'Código do produto', 1: 'Nome do produto', 2: 'Preço de Tabela', 16: 'Categoria principal', 17: 'Subcategoria nível 2', 18: 'Subcategoria nível 3', 19: 'Ativo / Inativo', 20: 'Exibido / Não exibido no e-commerce'}


def read_rows(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = book['Planilha1']
        headers = next(ws.iter_rows(values_only=True))
        if any(str(headers[i]).split('\n')[0] != expected for i, expected in HEADERS.items()):
            raise ValueError('Cabeçalho Mercos mudou; importação interrompida.')
        return [{**{key: row[i] for key, i in COLUMNS.items()}, 'sourceRow': number}
                for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2)
                if row[0] is not None or row[1] is not None]
    finally:
        book.close()


def prepare(rows, previous):
    key = lambda value: sheet.clean(value).upper()
    counts = collections.Counter(key(r['sku']) for r in rows if key(r['sku']))
    prior = {key(p['sku']): p for p in previous.get('products', [])}
    products, excluded = [], []
    changed = retained = 0
    for row in rows:
        sku, name = sheet.clean(row['sku']), sheet.clean(row['name'])
        price = sheet.cents(row['price'])
        reasons = []
        if not sku:
            reasons.append('Código ausente')
        elif counts[key(sku)] != 1:
            reasons.append('Código duplicado')
        if row['inactive'] != 0:
            reasons.append('Inativo ou situação não confirmada')
        if row['hidden'] != 0:
            reasons.append('Não exibido ou exibição não confirmada')
        if price is None or not name:
            reasons.append('Preço padrão ou nome inválido')
        if reasons:
            excluded.append({'sourceRow': row['sourceRow'], 'sku': sku, 'reasons': reasons})
            continue
        old = prior.get(key(sku))
        group = sheet.clean(row['group'])
        normalized_group = group.upper().replace('DO•HA', 'DOHA')
        brand = normalized_group if normalized_group in BRANDS else 'Marca a confirmar'
        # A prior explicitly supplied brand is usable when the Mercos group is empty.
        if not group and old and old['brand'] in BRANDS:
            brand = old['brand']
        size = sheet.size_from_name(name) or 'Embalagem a confirmar'
        category = sheet.clean(row['category']) or 'Grupo não informado'
        pending = ['Foto oficial', 'Taxonomia final', 'Disponibilidade comercial']
        if brand == 'Marca a confirmar':
            pending.append('Marca')
        if size == 'Embalagem a confirmar':
            pending.append('Embalagem')
        product = {
            'id': old['id'] if old else sku, 'sku': sku,
            'slug': old['slug'] if old else sheet.slugify(f'{name}-{sku}'),
            'name': name, 'brand': brand, 'sourceGroup': group,
            'line': category, 'category': category, 'size': size,
            'priceCents': price, 'status': 'consult',
            'order': old['order'] if old else len(prior) + row['sourceRow'],
            'keywords': [sku, group, category, sheet.clean(row['subcategory'])],
            'reviewPending': pending,
        }
        if old and old.get('image'):
            for field in ['image', 'imageAlt', 'imageSource']:
                if field in old:
                    product[field] = old[field]
            pending.remove('Foto oficial')
            pending.append('Embalagem da foto oficial')
        products.append(product)
        if old:
            retained += 1
            changed += old['priceCents'] != price
    products.sort(key=lambda p: p['order'])
    live_ids = {p['id'] for p in products}
    audit = {'excluded': excluded, 'previousMissing': [p['sku'] for p in prior.values() if p['id'] not in live_ids]}
    summary = {
        'sourceRows': len(rows), 'previewProducts': len(products),
        'excludedRows': len(excluded), 'retainedProducts': retained,
        'newProducts': len(products) - retained, 'changedPrices': changed,
        'previousProductsRemoved': len(audit['previousMissing']),
        'priceField': 'Mercos — Preço de Tabela', 'priceDate': '2026-10-07',
        'commercialOrdersEnabled': False, 'officialImageCandidates': sum(bool(p.get('image')) for p in products),
    }
    return {'demo': False, 'reviewOnly': True, 'taxonomyMode': 'mercos-group', 'summary': summary, 'products': products}, audit


def main():
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    source, prior_file, output = map(Path, sys.argv[1:])
    payload, audit = prepare(read_rows(source), json.loads(prior_file.read_text()))
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')
    # Local reconciliation stays next to the private input, never under public/.
    (source.parent / 'mercos-import-review.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(payload['summary'], ensure_ascii=False))


if __name__ == '__main__':
    main()
